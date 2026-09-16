"""Source and fixture checks; these do not claim installed EEM acceptance."""

import csv
import hashlib
import io
import json
import shutil
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

from generators.northlake_packet import GENERATED_DOCUMENTS, ROOT, SAMPLES, build_packet, render_generated_document, sample_files
from generators.update_workbook import TARGET, update_workbook


def packet_from_edited_sources(edit_file, old, new):
    """Rebuild the packet from copied sources with one text change applied."""
    with tempfile.TemporaryDirectory(prefix="northlake-source-test-") as folder:
        root = Path(folder)
        shutil.copytree(ROOT / "data", root / "data")
        file = root / edit_file
        text = file.read_text(encoding="utf-8")
        assert text.count(old) == 1, old
        file.write_text(text.replace(old, new), encoding="utf-8")
        return build_packet(root)


class NorthlakePacketTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.packet = build_packet()
        cls.sheets = cls.packet["sheets"]

    def test_counts_and_table_shapes_come_from_one_join(self):
        expected = {"companies": 3, "sites": 3, "buildings": 13, "providers": 5, "accounts": 49, "utilityMeters": 55, "ownedMeters": 265, "weatherStations": 2,
                    "meters": 322, "points": 431, "pointsToEnter": 299, "pointsGeneratedByEem": 264, "apartments": 124, "relatedMeasurements": 139,
                    "aggregates": 7, "gatewayProfiles": 15, "gatewayFormats": 21, "eventPublishers": 1, "baseRates": 14, "contacts": 12, "organizationUnits": 6}
        self.assertEqual(expected, self.packet["counts"])
        for name, sheet in self.sheets.items():
            for row in sheet["rows"]:
                self.assertEqual(len(sheet["headers"]), len(row), (name, row))
        text = json.dumps(self.packet, default=str)
        for stale in ["Applied Science", "Main Library", "Organization Units", "Portfolio", "Service Area", "Served Building", "supportLocations"]:
            self.assertNotIn(stale, text)

    def test_each_building_appears_once_in_the_company_that_owns_its_site(self):
        rows = self.sheets["Hierarchy"]["rows"]
        hierarchy = {f"{row[5]} / {row[1]}": row for row in rows}
        self.assertEqual(len(rows), len(hierarchy))
        self.assertEqual(3, sum(row[4] == "Company" for row in rows))
        self.assertEqual({"Northlake Main Campus", "Weather Reference", "Cedar Row Apartments", "Northlake Central Plant"}, {row[1] for row in rows if row[4] == "Site"})
        buildings = [row for row in rows if row[4] == "Location"]
        self.assertEqual(13, len(buildings))
        self.assertEqual(13, len({row[1] for row in buildings}))
        self.assertTrue(all(row[3] == "Building" for row in buildings))
        self.assertEqual(["Central Plant"], [row[1] for row in buildings if row[0] == "Northlake Thermal Plant"])
        self.assertTrue({"Science Center", "Library"} <= {row[1] for row in buildings})
        for path, row in hierarchy.items():
            self.assertTrue(row[5] == "System" or row[5] in hierarchy, path)
        meters = {row[0]: row for row in self.sheets["Meters"]["rows"]}
        self.assertEqual(self.packet["counts"]["meters"], len(meters))
        for row in meters.values():
            self.assertIn(row[-1], hierarchy)
            self.assertEqual(row[7], hierarchy[row[-1]][0])
        served = [row for row in meters.values() if row[4] == "Science Center" and row[2] == "Chilled Water"]
        self.assertEqual(1, len(served))
        self.assertEqual(["Northlake Thermal Plant", "Utility meter", "Northlake University"], served[0][5:8])
        self.assertEqual("System / Northlake University / Northlake Main Campus / Science Center", served[0][-1])
        plant = [row for row in meters.values() if row[7] == "Northlake Thermal Plant"]
        self.assertEqual(8, len(plant))
        self.assertTrue(all(row[4] == "Central Plant" for row in plant))
        self.assertEqual({"Utility meter", "Production meter", "Plant controller", "Sewer-deduct submeter"}, {row[6] for row in plant})
        self.assertTrue(all(row[1] in meters for row in self.sheets["Measured Points"]["rows"]))

    def test_billing_accounts_carry_their_editor_setup(self):
        accounts = {row[3]: row for row in self.sheets["Accounts And Agreements"]["rows"]}
        details = {row[0]: row for row in self.sheets["Account Billing Details"]["rows"]}
        self.assertEqual(49, len(accounts))
        self.assertEqual(set(accounts), set(details))
        for number, row in details.items():
            self.assertTrue(row[2].endswith(", CA 95822") and row[3] and row[4] and "@" in row[5], number)
        self.assertEqual("100 Synthetic Campus Drive, Sacramento, CA 95822", details["SYN-VED-A-0010001"][2])
        self.assertEqual("PO Box 41027, Sacramento, CA 95841", details["SYN-VED-A-0010001"][3])
        self.assertEqual("Owen Blake", details["SYN-VED-A-0010001"][4])
        self.assertEqual("Devon Brooks", details["SYN-NTP-C-0000001"][4])

        setup = {row[0]: row for row in self.packet["accountSetup"]["rows"]}
        self.assertEqual(set(accounts), set(setup))
        self.assertEqual({"AP", "AP and GL", "GL", "No Upload"}, {row[4] for row in setup.values()})
        self.assertEqual({"Bill Importer", "Bill Entry", "Bill Allocation"}, {row[3] for row in setup.values()})
        self.assertEqual("Bill Importer", setup["SYN-VED-A-0010001"][3])          # campus electric arrives as a file
        self.assertEqual("Bill Entry", setup["SYN-VED-A-0010005"][3])             # Cedar Row is keyed by hand
        self.assertEqual("Bill Allocation", setup["SYN-NTP-C-0000001"][3])
        self.assertEqual("No Upload", setup["SYN-RCU-A-0030004"][4])              # the never-exported negative control
        self.assertTrue(all(row[5] == "Energy Hippo Default Template" for row in setup.values()))
        self.assertTrue(all(row[6] == "2000-AP" for row in setup.values()))
        self.assertEqual({"5100-ELEC", "5110-GAS", "5120-WATER", "5130-THERMAL", "5140-SOLAR"}, {row[7] for row in setup.values()})

        overrides = self.packet["validationOverrides"]["rows"]
        self.assertEqual(4, len(overrides))
        self.assertEqual({"On", "Off"}, {row[3] for row in overrides})
        self.assertTrue(all(row[0] in accounts and row[5] for row in overrides))
        self.assertEqual(45, sum(row[8] == "Inherits company" for row in setup.values()))

    def test_bill_only_meters_have_no_points(self):
        points = {}
        for row in self.sheets["Measured Points"]["rows"]:
            points.setdefault(row[1], []).append(row[0])
        meters = {row[0]: row for row in self.sheets["Meters"]["rows"]}
        self.assertNotIn("Admin Hall Electric Main", points)
        self.assertEqual("Bill / allocation statement", meters["Admin Hall Electric Main"][8])
        self.assertEqual(["Delivered kWh 15m", "Demand kW 15m"], points["Student Center Electric Main"])
        steam = [name for name, row in meters.items() if row[2] == "Steam" and row[7] == "Northlake University"]
        self.assertEqual(6, len(steam))
        self.assertFalse(set(steam) & set(points))
        self.assertIn("Central Plant Steam Production", points)
        self.assertEqual(sum(len(names) for names in points.values()), self.packet["counts"]["points"])

    def test_cedar_row_is_master_metered_with_owned_unit_submeters(self):
        units = self.sheets["Units and Submeters"]["rows"]
        self.assertEqual(124, len(units))
        by_building = {}
        for row in units:
            by_building.setdefault(row[0], []).append(row)
        self.assertEqual([60, 64], [len(by_building["Cedar Row A"]), len(by_building["Cedar Row B"])])
        self.assertEqual(["101", "320", "416"], [by_building["Cedar Row A"][0][1], by_building["Cedar Row A"][-1][1], by_building["Cedar Row B"][-1][1]])
        for row in units:
            letter = row[0][-1]
            self.assertEqual([f"SYN-SUB-E-{letter}{row[1]}", "Submeter kWh 60m", f"SYN-SUB-W-{letter}{row[1]}", "Register Reading", "Route Usage kgal"], row[3:8])
        meters = {row[0]: row for row in self.sheets["Meters"]["rows"]}
        unit_meter = meters["Cedar Row A Unit 214 Electric"]
        self.assertEqual(["Apartment submeter", "Cedar Row Apartments"], unit_meter[6:8])
        self.assertIn("Behind Cedar Row A Electric Master", unit_meter[10])
        self.assertIn("No billing account", unit_meter[10])
        rollups = {}
        for row in self.sheets["Rollup Members"]["rows"]:
            rollups.setdefault(row[0], []).append(row)
        self.assertEqual(7, len(rollups))
        self.assertEqual([61, 65], [len(rollups["Cedar Row A Submetered kWh Total"]), len(rollups["Cedar Row B Submetered kWh Total"])])
        sources = self.sheets["Source Measurements"]["rows"]
        self.assertEqual(128, sum(row[0] == "Fixed-network electric submeters" for row in sources))
        self.assertEqual(126, sum(row[0] == "Cedar Row apartment water submeter route" for row in sources))
        self.assertEqual(len(sources), len({(row[0], row[3], row[4]) for row in sources}))

    def test_relationships_and_format_cases_are_explicit(self):
        relationships = self.sheets["Related Measurements"]["rows"]
        self.assertEqual(132, sum(row[4] == "Usage from register difference" for row in relationships))
        self.assertEqual(4, sum("degree days" in row[4] for row in relationships))
        self.assertEqual(3, sum(row[4] == "Measured versus baseline" for row in relationships))
        assignments = self.sheets["Weather Assignments"]["rows"]
        self.assertEqual(13, len(assignments))
        self.assertTrue(all(row[2] == "KSAC" for row in assignments))
        coverage = self.packet["coverage"]["rows"]
        formats = {row[1] for row in coverage}
        self.assertTrue({"ACQUISUITE", "MDEF", "MV9", "BACNET", "MODBUS", "FIG", "SIEMANSREPORT", "MEDIATOR", "EATON", "XML", "FIXEDNETWORKS", "CMEP", "NEPTUNE", "MVRS", "HMR_EVENTS"} <= formats)
        self.assertTrue(all(row[-2:] == ["Not run", "Not run"] for row in coverage))
        self.assertEqual(["Generated"], [row[4] for row in coverage if row[0] == "gw_acquisuite_student_electric"])
        records = self.packet["instanceMapping"]["rows"]
        self.assertEqual(322 + 431 + 4 + 3 + 7, len(records))  # meters + points + degree-day + baseline + aggregate points.
        self.assertEqual(len(records), len({row[1] for row in records}))
        self.assertTrue(all(row[19:23] == ["", "", "", ""] for row in records))   # the installed IDs stay blank

    def test_meters_and_points_carry_their_editor_setup(self):
        meters = {row[0]: row for row in self.packet["meterSetup"]["rows"]}
        self.assertEqual(322, len(meters))
        self.assertEqual(55, sum(1 for row in meters.values() if row[7] != ""))          # a bill cycle needs a provider
        self.assertEqual(44, sum(1 for row in meters.values() if row[4]))                # provider badge numbers
        self.assertEqual("VED-BADGE-0040001", meters["SYN-VED-M-0040001"][4])
        self.assertEqual(1, meters["SYN-VED-M-0040001"][7])
        self.assertEqual(2, meters["SYN-RCU-M-0060003"][7])                              # Cedar Row runs the second cycle
        self.assertEqual([40, 10], [row[5] for row in meters.values() if row[5]])

        points = self.packet["pointSetup"]["rows"]
        entered = sum(row[6] for row in points if row[5].startswith("Enter"))
        generated = sum(row[6] for row in points if row[5].startswith("Generated"))
        self.assertEqual([299, 132], [entered, generated])
        self.assertEqual({5, 15, 60, "None"}, {row[3] for row in points})             # only intervals the editor offers
        self.assertTrue(all(row[1] == "Digital" and row[2] == "State" for row in points if row[0] == "Digital"))
        self.assertTrue(all(row[1] == "Temperature" for row in points if row[4] != "None"))
        self.assertEqual(2, sum(row[6] for row in points if row[4] != "None"))           # the two degree-day drivers

        providers = {row[0]: row for row in self.packet["providerSetup"]["rows"]}
        self.assertEqual(5, len(providers))
        self.assertTrue(all(row[3] and row[4] and row[5] for row in providers.values()))  # remit, cycles, vendor numbers
        self.assertIn("2 Cedar Row", providers["River City Utilities"][4])
        self.assertEqual("None - internal counterparty", providers["Northlake Thermal Plant"][1])

    def test_tariff_and_area_values_come_from_existing_sources(self):
        buildings = {row[0]: row for row in self.sheets["Buildings"]["rows"]}
        self.assertEqual([95000, 58000, 110000, ""], [buildings["Student Center"][5], buildings["Cedar Row A"][5], buildings["Lakeview Residence Hall"][5], buildings["Campus Grounds"][5]])
        schedules = {row[1] for row in self.sheets["Rate Schedules"]["rows"]}
        self.assertTrue({"VED-TOU-GS", "VED-TOU-GS-FY25", "VED-TOU-PRI", "VED-AL-1", "SGU-GL-1", "RCU-IRR", "RCU-FIRE", "NTP-ALLOC-CHW"} <= schedules)
        charges = self.sheets["Rate Components"]["rows"]
        peak = next(row for row in charges if row[1] == "VED-TOU-GS" and row[2] == "Energy / peak")
        self.assertEqual(0.241, peak[3])
        self.assertTrue(all(isinstance(row[3], (int, float)) for row in charges))
        accounts = self.sheets["Accounts And Agreements"]["rows"]
        self.assertEqual(49, len({row[3] for row in accounts}))
        self.assertEqual({"External utility account", "Internal cost center", "PPA agreement"}, {row[2] for row in accounts})

    def test_drifted_identifiers_and_ownership_are_rejected(self):
        with self.assertRaisesRegex(AssertionError, "Account identifier drift: ved_student_center"):
            packet_from_edited_sources("data/scenarios/demo-university-v1.yaml", "SYN-VED-A-0010004", "WRONG-ACCOUNT")
        with self.assertRaisesRegex(AssertionError, "Stage 1 sites differ from scenario ownership: northlake_thermal_plant"):
            packet_from_edited_sources("data/eem/northlake-eem-stage1-setup-v1.yaml", "hierarchySitesFromScenario: [northlake_central_plant]", "hierarchySitesFromScenario: [northlake_central_plant, northlake_main_campus]")
        with self.assertRaisesRegex(AssertionError, "Aggregate parent is not a hierarchy node: agg_nlu_science_net_electric_kwh_15m"):
            packet_from_edited_sources("data/eem/northlake-eem-metaworld-stage1b-v1.yaml", "parentPath: System / Northlake University / Northlake Main Campus / Science Center", "parentPath: System / Northlake University / Northlake Main Campus / Applied Science Center")
        with self.assertRaisesRegex(AssertionError, "Duplicate building path"):
            packet_from_edited_sources("data/scenarios/demo-university-v1.yaml", "        displayName: Library\n", "        displayName: Science Center\n")
        with self.assertRaisesRegex(AssertionError, "Security scope buildings differ from scenario ownership: northlake_thermal_plant"):
            packet_from_edited_sources("data/security/northlake-eem-security-v1.yaml", "      buildings: [central_plant]", "      buildings: [central_plant, science_center]")

    def test_company_contacts_cover_the_handover_and_department_ownership(self):
        contacts = {row[1]: row for row in self.sheets["Contacts"]["rows"]}
        self.assertEqual(12, len(contacts))
        self.assertTrue({"Samantha Ireland", "Jordan Hale", "Priya Nandakumar", "Iris Morales", "Nora Chen", "Leo Martinez", "Maya Chen", "Devon Brooks"} <= contacts.keys())
        self.assertEqual("samantha.ireland@northlake.example.edu", contacts["Samantha Ireland"][3])
        self.assertEqual("Cedar Row Apartments", contacts["Maya Chen"][0])
        self.assertEqual("Northlake Thermal Plant", contacts["Devon Brooks"][0])
        self.assertEqual("facilities.ops@northlake.example.edu", contacts["Jordan Hale"][3])
        self.assertTrue(all(row[3] and row[4] and "TBD" not in row for row in contacts.values()))
        self.assertTrue({"Valley Electric District", "Sierra Gas Utility", "River City Utilities", "Helios Onsite Solar", "Northlake Thermal Plant"} <= {row[0] for row in contacts.values()})
        units = self.sheets["Departments & Responsibilities"]["rows"]
        self.assertEqual(6, len(units))
        self.assertEqual({"Northlake University", "Cedar Row Apartments", "Northlake Thermal Plant"}, {row[0] for row in units})
        for row in units:
            contact = contacts[row[3]]
            self.assertEqual([contact[0], contact[3], contact[4]], [row[0], row[4], row[5]])
        departments = {row[1] for row in units}
        self.assertTrue(all(row[6] in departments for row in self.sheets["Sites"]["rows"]))
        self.assertTrue(all(row[3] in departments for row in self.sheets["Buildings"]["rows"]))

    def test_acquisuite_golden_delivery_has_exact_time_and_energy_contract(self):
        files = sample_files()
        rows = list(csv.reader(io.StringIO(files["complete"]["content"])))
        self.assertEqual(96, len(rows))
        times = [datetime.strptime(row[0], "'%Y-%m-%d %H:%M:%S'") for row in rows]
        self.assertEqual(datetime(2024, 1, 15, 8), times[0])
        self.assertEqual(datetime(2024, 1, 16, 7, 45), times[-1])
        self.assertTrue(all(b - a == timedelta(minutes=15) for a, b in zip(times, times[1:])))
        self.assertEqual(2020, sum(float(row[1]) for row in rows))
        self.assertEqual(120, max(float(row[2]) for row in rows))
        self.assertTrue(all(float(row[1]) == float(row[2]) / 4 for row in rows))
        self.assertEqual("ASQVED01_STUDENT.log", files["complete"]["filename"])

    def test_gap_and_backfill_are_disjoint_and_reconstruct_the_complete_file(self):
        files = sample_files()
        complete = files["complete"]["content"].splitlines()
        gap = files["gap"]["content"].splitlines()
        backfill = files["backfill"]["content"].splitlines()
        self.assertEqual(88, len(gap))
        self.assertEqual(8, len(backfill))
        self.assertFalse(set(gap) & set(backfill))
        self.assertEqual(complete, sorted(gap + backfill))
        self.assertTrue(backfill[0].startswith("'2024-01-15 20:00:00'"))
        self.assertTrue(backfill[-1].startswith("'2024-01-15 21:45:00'"))

    def test_checked_in_documents_match_regeneration(self):
        for pdf in GENERATED_DOCUMENTS:
            rendered = io.BytesIO()
            render_generated_document(self.packet, pdf, rendered)
            self.assertTrue((ROOT / pdf).read_bytes() == rendered.getvalue(), f"{pdf} is stale; run python generators/northlake_packet.py")

    def test_checked_in_deliveries_and_hashes_match_regeneration(self):
        directory = ROOT / SAMPLES
        expected = json.loads((directory / "expected-results.json").read_text(encoding="utf-8"))
        self.assertEqual("Not run", expected["installedAcceptance"])
        for case, sample in sample_files().items():
            actual = (directory / case / sample["filename"]).read_bytes()
            self.assertEqual(sample["content"].encode(), actual)
            self.assertEqual(expected["cases"][case]["sha256"], hashlib.sha256(actual).hexdigest())

    def test_workbook_updater_rewrites_generated_sheets_and_preserves_the_rest(self):
        with tempfile.TemporaryDirectory(prefix="northlake-workbook-test-") as folder:
            copy = Path(folder) / "workbook.xlsx"
            shutil.copyfile(TARGET, copy)
            update_workbook(target=copy, packet=self.packet)
            self.assertEqual([], update_workbook(target=copy, packet=self.packet))
            wb = openpyxl.load_workbook(copy)
            self.assertEqual(["Instructions", "Hierarchy"], wb.sheetnames[:2])
            self.assertTrue({"Service Profiles", "Units and Submeters", "Chart of Accounts", "Users & Access", "Tenants & Leases", "Sustainability", "Projects"} <= set(wb.sheetnames))
            self.assertNotIn("Organization Units", wb.sheetnames)
            self.assertEqual(len(self.sheets) + 5, len(wb.sheetnames))
            meters = wb["Meters"]
            self.assertEqual("Meter Name", meters["A4"].value)
            self.assertEqual(self.packet["counts"]["meters"] + 4, meters.max_row)
            self.assertEqual("B5", meters.freeze_panes)
            self.assertEqual(f"A4:L{meters.max_row}", meters.tables["MetersTable"].ref)
            self.assertEqual("Town Center", wb["Tenants & Leases"]["A5"].value)


if __name__ == "__main__":
    unittest.main()
